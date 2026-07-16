"""Evaluate the Excel/Krchnak difficult-coupling predictor over both classes.

Extracts per-residue aggregation potentials from a scored 'Difficulty (Sheet)' and
scores each peptide as a synthesis-outcome predictor, so the Excel tool gets an AUC
directly comparable to peptimizer (0.71) and PepFuNN (0.55).

Tool: Peptide Companion (M. Lebl), implementing Krchnak, Flegelova & Vagner,
Int. J. Peptide Protein Res. 42, 450 (1993). The macro writes one cell per residue,
"<AA> <coefficient>", and fills difficult residues red. Higher coefficient = more
aggregation = harder coupling.

Workflow (the macro must run in Excel first - it cannot be driven from here):
  1. Open experiments/predictors/Prediction_all_peptides.xlsm (all 1771 peptides
     already in column A of 'Difficulty (Sheet)').
  2. Click the "Difficult couplings predictor" button, save, close.
  3. Run from the project root:
       python experiments/predictors/excel_difficulty_evaluate.py

Per-peptide scores:
  max_coeff  = worst residue's potential (headline; analogue of peptimizer's worst step)
  mean_coeff = average potential
  red_cells  = number of residues the macro flagged red (the tool's own verdict)
"""
import collections
import csv
import datetime
import os
import re

import openpyxl

PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
XLSM = os.path.join(PROJECT_ROOT, "experiments", "predictors", "Prediction_all_peptides.xlsm")
CSV_SRC = os.path.join(PROJECT_ROOT, "data", "raw", "peptide_baza.csv")
OUT_DIR = os.path.join(PROJECT_ROOT, "experiments", "predictors")
SHEET = "Difficulty (Sheet)"
RED, GREEN = "FFFF0000", "FF00FF00"


def flags_by_sequence(path=CSV_SRC):
    with open(path, newline="", encoding="utf-8") as fh:
        out = {}
        for r in csv.DictReader(fh, delimiter=";"):
            f = r["synthesis_flag"].strip().upper()
            if f in ("TRUE", "FALSE"):
                out[r["peptide_seq"].strip()] = f
    return out


def parse_coeff(text):
    """'L 1,32' -> 1.32 ; '0.88' default uses a dot, computed values use a comma."""
    tok = str(text).split()[-1].replace(",", ".")
    return float(tok)


def cell_color(c):
    fg = c.fill.fgColor.rgb if (c.fill and c.fill.patternType) else None
    fc = c.font.color.rgb if (c.font and c.font.color and c.font.color.type == "rgb") else None
    for v in (fg, fc):
        if v == RED:
            return "RED"
        if v == GREEN:
            return "GREEN"
    return None


def extract(ws):
    rows = []
    for r in range(2, ws.max_row + 1):
        seq = ws.cell(r, 1).value
        if not seq:
            continue
        coeffs, reds, greens = [], 0, 0
        for cc in range(2, ws.max_column + 1):
            c = ws.cell(r, cc)
            if c.value is None:
                continue
            try:
                coeffs.append(parse_coeff(c.value))
            except (ValueError, IndexError):
                continue
            col = cell_color(c)
            reds += col == "RED"
            greens += col == "GREEN"
        if not coeffs:
            continue
        rows.append({"sequence": str(seq).strip(), "length": len(str(seq).strip()),
                     "n_residues": len(coeffs),
                     "max_coeff": round(max(coeffs), 4),
                     "mean_coeff": round(sum(coeffs) / len(coeffs), 4),
                     "red_cells": reds, "green_cells": greens})
    return rows


def auc_roc(scores, labels):
    pos = [s for s, y in zip(scores, labels) if y == 1]
    neg = [s for s, y in zip(scores, labels) if y == 0]
    if not pos or not neg:
        return float("nan")
    order = sorted(range(len(scores)), key=lambda i: scores[i])
    ranks = [0.0] * len(scores)
    i = 0
    while i < len(order):
        j = i
        while j + 1 < len(order) and scores[order[j + 1]] == scores[order[i]]:
            j += 1
        avg = (i + j) / 2 + 1
        for k in range(i, j + 1):
            ranks[order[k]] = avg
        i = j + 1
    rank_pos = sum(r for r, y in zip(ranks, labels) if y == 1)
    return (rank_pos - len(pos) * (len(pos) + 1) / 2) / (len(pos) * len(neg))


def main():
    if not os.path.exists(XLSM):
        raise SystemExit(f"not found: {XLSM}\nRun the macro on all 1771 peptides first (see docstring).")

    flags = flags_by_sequence()
    ws = openpyxl.load_workbook(XLSM, keep_vba=True)[SHEET]
    rows = extract(ws)

    # detect "not scored yet": every residue still at default 0.88 across the board
    all_default = all(abs(r["max_coeff"] - 0.88) < 1e-9 for r in rows)
    if all_default:
        raise SystemExit("Sheet appears UNSCORED (all coefficients = 0.88). "
                         "Open the workbook, click the Predict button, save, then re-run.")

    for r in rows:
        r["synthesis_flag"] = flags.get(r["sequence"], "?")
    missing = [r for r in rows if r["synthesis_flag"] == "?"]
    if missing:
        print(f"WARNING: {len(missing)} sheet sequences not found in the CSV (skipped in AUC)")

    stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
    detail = os.path.join(OUT_DIR, f"excel_difficulty_report_{stamp}.csv")
    with open(detail, "w", newline="", encoding="utf-8") as fh:
        cols = ["sequence", "synthesis_flag", "length", "n_residues",
                "max_coeff", "mean_coeff", "red_cells", "green_cells"]
        w = csv.DictWriter(fh, fieldnames=cols, delimiter=";")
        w.writeheader()
        w.writerows({k: r[k] for k in cols} for r in rows)

    scored = [r for r in rows if r["synthesis_flag"] in ("TRUE", "FALSE")]
    F = [r for r in scored if r["synthesis_flag"] == "FALSE"]
    T = [r for r in scored if r["synthesis_flag"] == "TRUE"]

    L = ["Excel / Krchnak difficult-coupling predictor - both classes",
         f"Generated: {datetime.datetime.now():%Y-%m-%d %H:%M:%S}",
         "Peptide Companion (M. Lebl); Krchnak et al., Int. J. Peptide Protein Res. 42, 450 (1993)",
         f"Source workbook: {os.path.basename(XLSM)}",
         f"Peptides scored: {len(scored)}  ({len(T)} TRUE / {len(F)} FALSE)",
         "",
         "Higher coefficient = more aggregation = harder coupling.",
         ""]

    if F and T:
        import statistics as st

        def line(key, rows_):
            v = [r[key] for r in rows_]
            return f"mean {st.mean(v):+.4f}  sd {st.pstdev(v):.4f}  min {min(v):+.4f}  max {max(v):+.4f}"

        for key in ("max_coeff", "mean_coeff", "red_cells"):
            L.append(f"== {key} ==")
            L.append(f"  FALSE: {line(key, F)}")
            L.append(f"  TRUE : {line(key, T)}")
            L.append("")

        labels = [1 if r["synthesis_flag"] == "FALSE" else 0 for r in scored]
        auc_max = auc_roc([r["max_coeff"] for r in scored], labels)
        auc_mean = auc_roc([r["mean_coeff"] for r in scored], labels)
        auc_red = auc_roc([r["red_cells"] for r in scored], labels)
        auc_len = auc_roc([r["length"] for r in scored], labels)

        s15 = [r for r in scored if r["length"] == 15]
        lab15 = [1 if r["synthesis_flag"] == "FALSE" else 0 for r in s15]
        auc_max15 = auc_roc([r["max_coeff"] for r in s15], lab15)
        auc_red15 = auc_roc([r["red_cells"] for r in s15], lab15)

        L += ["== Discrimination (positive = FALSE / will fail) ==",
              f"  AUC-ROC, max_coeff   : {auc_max:.3f}   (0.500 = chance)",
              f"  AUC-ROC, mean_coeff  : {auc_mean:.3f}",
              f"  AUC-ROC, red_cells   : {auc_red:.3f}   (the tool's own verdict)",
              f"  AUC-ROC, length only : {auc_len:.3f}   (confound)",
              "",
              "== Length control (15-mer stratum) ==",
              f"  n={len(s15)} ({sum(lab15)} FALSE / {len(s15)-sum(lab15)} TRUE)",
              f"  AUC-ROC, max_coeff   : {auc_max15:.3f}",
              f"  AUC-ROC, red_cells   : {auc_red15:.3f}",
              ""]

    L.append(f"Per-peptide detail: {os.path.basename(detail)}")
    summary = os.path.join(OUT_DIR, f"excel_difficulty_report_{stamp}.txt")
    with open(summary, "w", encoding="utf-8") as fh:
        fh.write("\n".join(L) + "\n")
    print("\n".join(L))
    print()
    print("wrote:", detail)
    print("wrote:", summary)


if __name__ == "__main__":
    main()
