"""Report red/green/uncolored difficulty categories from Prediction.xlsm.

Run from the project root, after opening Prediction.xlsm and clicking the
"Difficult couplings predictor" button on the "Difficulty (Sheet)" tab:

    python experiments/predictors/difficulty_report.py

Read-only with respect to the workbook - it only writes the two report files.

Background: the workbook is Peptide Companion (M. Lebl), implementing Krchnak et al.,
Int. J. Peptide Protein Res. 42, 450 (1993). Its macro scores each residue's
aggregation potential and marks the per-residue score cells by FILL colour:
red FFFF0000 = difficult coupling, green FF00FF00 = easy. Column A rows 2..N hold
the sequences (the 4 bundled demo peptides have been removed).

Category rule: a sequence with ANY red cell -> RED (even if it also has green
cells); else any green cell -> GREEN; else NONE.
"""
import collections
import csv
import datetime
import os

import openpyxl

XLSM = os.path.join("experiments", "predictors", "Prediction.xlsm")
CSV_SRC = os.path.join("data", "raw", "peptide_baza.csv")
OUT_DIR = os.path.join("experiments", "predictors")
SHEET = "Difficulty (Sheet)"
RED, GREEN = "FFFF0000", "FF00FF00"


def read_false_sequences(path=CSV_SRC):
    with open(path, newline="", encoding="utf-8") as fh:
        return [r["peptide_seq"].strip()
                for r in csv.DictReader(fh, delimiter=";")
                if r["synthesis_flag"].strip().upper() == "FALSE"]


def cell_color(c):
    """Return 'RED'/'GREEN'/None for a scored residue cell.

    Checks fill and font: the macro uses fill, but styles for coloured fonts also
    exist in the workbook, so both are inspected rather than assumed.
    """
    if c.value is None:
        return None
    fg = c.fill.fgColor.rgb if (c.fill and c.fill.patternType) else None
    fc = c.font.color.rgb if (c.font and c.font.color and c.font.color.type == "rgb") else None
    for v in (fg, fc):
        if v == RED:
            return "RED"
        if v == GREEN:
            return "GREEN"
    return None


def main():
    src = read_false_sequences()
    ws = openpyxl.load_workbook(XLSM, keep_vba=True)[SHEET]

    # The sheet must hold exactly the FALSE peptides, in CSV order, from row 2.
    got = [ws.cell(r, 1).value for r in range(2, 2 + len(src))]
    if got != src:
        raise SystemExit(f"{SHEET} rows 2.. do not match the CSV FALSE list - re-check the workbook")

    rows = []
    for i, seq in enumerate(src):
        r = 2 + i
        counts = collections.Counter(
            cell_color(ws.cell(r, col)) for col in range(2, ws.max_column + 1)
        )
        red, green = counts["RED"], counts["GREEN"]
        rows.append({"row": r, "sequence": seq, "length": len(seq),
                     "category": "RED" if red else ("GREEN" if green else "NONE"),
                     "red_cells": red, "green_cells": green})

    cat = collections.Counter(x["category"] for x in rows)
    both = sum(1 for x in rows if x["red_cells"] and x["green_cells"])
    n = len(rows)

    stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
    detail = os.path.join(OUT_DIR, f"difficulty_report_{stamp}.csv")
    with open(detail, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=["row", "sequence", "length", "category",
                                           "red_cells", "green_cells"], delimiter=";")
        w.writeheader()
        w.writerows(rows)

    lines = [
        f"Difficult-coupling prediction report - Prediction.xlsm / '{SHEET}'",
        f"Generated: {datetime.datetime.now():%Y-%m-%d %H:%M:%S}",
        f"Source: {CSV_SRC}  (synthesis_flag = FALSE)",
        "",
        f"Sequences analysed: {n}   (4 demo peptides removed; not counted)",
        "",
        "Category rule: any red cell -> RED (even if green cells also present);",
        "               else any green cell -> GREEN; else NONE.",
        "",
        f"  RED   (>=1 red cell)     : {cat['RED']:3d}  ({cat['RED']/n*100:5.1f}%)",
        f"  GREEN (green, no red)    : {cat['GREEN']:3d}  ({cat['GREEN']/n*100:5.1f}%)",
        f"  NONE  (no coloured cells): {cat['NONE']:3d}  ({cat['NONE']/n*100:5.1f}%)",
        f"  {'-'*40}",
        f"  TOTAL                    : {sum(cat.values()):3d}",
        "",
        f"Of the {cat['RED']} RED sequences, {both} also contain green cells",
        f"(counted as RED per the rule above); {cat['RED']-both} are red-only.",
        "",
        f"Per-sequence detail: {os.path.basename(detail)}",
    ]
    summary = os.path.join(OUT_DIR, f"difficulty_report_{stamp}.txt")
    with open(summary, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")

    print("\n".join(lines))
    print()
    print("wrote:", detail)
    print("wrote:", summary)


if __name__ == "__main__":
    main()
